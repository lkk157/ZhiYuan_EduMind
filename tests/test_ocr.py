# -*- coding: utf-8 -*-
"""
多模态入库单测（M3）：OCR 回填、★显存互斥时序、失败容错、流水线端到端。

为什么时序必须测死：「先卸 LLM → 再跑 OCR（keep_alive=0）」是 8G 显存不爆的命门，
它靠运行时行为保证、静态 review 看不出来——测试必须断言 unload 发生在 generate 之前，
且 OCR 调用的 keep_alive 恒为 0（用完即卸）。全程假网关，绝不打真实 Ollama。
"""
import base64
import io

import pytest

from app.core.config import settings
from app.core.exceptions import UpstreamError


class _FakeOcrGateway:
    """假 OCR 网关：记录调用时序，可配置驻留状态 / 固定回复 / 注入故障。"""

    def __init__(self, active: list[str] | None = None, reply: str = "识别文本 E=mc^2", fail: bool = False):
        self.active = list(active or [])
        self.reply = reply
        self.fail = fail
        self.calls: list[tuple] = []  # ("active",) / ("unload", model) / ("generate", kwargs)

    async def active_models(self) -> list[str]:
        self.calls.append(("active",))
        return list(self.active)

    async def unload(self, model: str) -> None:
        self.calls.append(("unload", model))

    async def generate(self, **kwargs) -> str:
        self.calls.append(("generate", kwargs))
        if self.fail:
            raise UpstreamError("OCR 服务不可用")
        return self.reply


def _png_bytes() -> bytes:
    """用 Pillow 现造一张小 PNG（10x10 纯色即可——OCR 走假网关，不真识别）。"""
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (10, 10), color=(200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _call_names(fake: _FakeOcrGateway) -> list[str]:
    return [c[0] for c in fake.calls]


def test_image_file_parsed_as_empty_page(tmp_path):
    """图片文件走 parse_document 分发：一页空文本 + empty=[1]，留给 OCR 回填。"""
    from app.ingest.parsers import parse_document

    path = tmp_path / "pic.png"
    path.write_bytes(_png_bytes())
    pages, empty = parse_document(path)
    assert [p.page_no for p in pages] == [1]
    assert pages[0].text == ""
    assert empty == [1]


# ===== 互斥序列与时序（本阶段红线命门） =====


@pytest.mark.asyncio
async def test_vram_mutex_sequence_unload_before_ocr(tmp_path, monkeypatch):
    """★ LLM 驻留时：unload 必须发生在 OCR generate 之前，且 keep_alive=0、模型名来自配置。"""
    from app.ingest.ocr import OCR_NUM_PREDICT, ocr_fill_pages
    from app.ingest.parsers import parse_document

    fake = _FakeOcrGateway(active=[settings.llm_model, settings.embed_model])
    monkeypatch.setattr("app.ingest.ocr.gateway", fake)

    path = tmp_path / "pic.png"
    path.write_bytes(_png_bytes())
    pages, empty = parse_document(path)
    failed = await ocr_fill_pages(path, pages, empty)

    names = _call_names(fake)
    # 时序断言：查驻留 → 卸 LLM → 再跑 OCR（顺序错了 8G 显存会同时装两类大模型）
    assert names.index("unload") < names.index("generate")
    assert ("unload", settings.llm_model) in fake.calls

    gen_kwargs = next(c[1] for c in fake.calls if c[0] == "generate")
    assert gen_kwargs["model"] == settings.ocr_model  # 模型名只从配置读
    assert gen_kwargs["keep_alive"] == settings.ocr_keep_alive  # "0" 用完即卸（红线）
    assert gen_kwargs["num_predict"] == OCR_NUM_PREDICT
    assert gen_kwargs["temperature"] == 0.0
    assert len(gen_kwargs["images"]) == 1
    # 图像必须是合法 base64（Ollama images 参数格式）
    assert len(base64.b64decode(gen_kwargs["images"][0])) > 0

    # 回填成功：empty 清单清空，文本写回页
    assert failed == []
    assert "E=mc^2" in pages[0].text


@pytest.mark.asyncio
async def test_no_unload_when_llm_not_resident(tmp_path, monkeypatch):
    """LLM 未驻留时绝不 unload（否则 Ollama 会先加载再卸载，白等几十秒）。"""
    from app.ingest.ocr import ocr_fill_pages
    from app.ingest.parsers import parse_document

    fake = _FakeOcrGateway(active=[])  # 显存里什么都没有
    monkeypatch.setattr("app.ingest.ocr.gateway", fake)

    path = tmp_path / "pic.png"
    path.write_bytes(_png_bytes())
    pages, empty = parse_document(path)
    failed = await ocr_fill_pages(path, pages, empty)

    assert "unload" not in _call_names(fake)  # 没驻留就不卸
    assert "generate" in _call_names(fake)  # OCR 照常执行
    assert failed == []


@pytest.mark.asyncio
async def test_no_ocr_touch_when_no_empty_pages(tmp_path, monkeypatch):
    """纯文本文档（empty_pages=[]）零模型调用：不查驻留、不卸、不识别。"""
    from app.ingest.ocr import ocr_fill_pages
    from app.ingest.parsers import ParsedPage

    fake = _FakeOcrGateway()
    monkeypatch.setattr("app.ingest.ocr.gateway", fake)
    failed = await ocr_fill_pages(tmp_path / "x.pdf", [ParsedPage(1, "有文本")], [])
    assert failed == []
    assert fake.calls == []


# ===== 失败容错 =====


@pytest.mark.asyncio
async def test_ocr_failure_keeps_empty_pages(tmp_path, monkeypatch):
    """OCR 故障不炸入库：页保留在 empty_pages、文本不写假内容，可重传重试。"""
    from app.ingest.ocr import ocr_fill_pages
    from app.ingest.parsers import parse_document

    fake = _FakeOcrGateway(active=[], fail=True)
    monkeypatch.setattr("app.ingest.ocr.gateway", fake)

    path = tmp_path / "pic.png"
    path.write_bytes(_png_bytes())
    pages, empty = parse_document(path)
    failed = await ocr_fill_pages(path, pages, empty)

    assert failed == [1]
    assert pages[0].text == ""  # 绝不写入半截/伪造内容


@pytest.mark.asyncio
async def test_blank_pdf_page_rendered_by_pymupdf(tmp_path, monkeypatch):
    """扫描版 PDF 空白页：pymupdf 渲染成图喂 OCR，回填后该页可切块（M2→M3 交接点）。"""
    from pypdf import PdfWriter

    from app.ingest.ocr import ocr_fill_pages
    from app.ingest.parsers import parse_document

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "scan.pdf"
    with path.open("wb") as f:
        writer.write(f)

    fake = _FakeOcrGateway(active=[], reply="扫描页识别出的正文内容")
    monkeypatch.setattr("app.ingest.ocr.gateway", fake)

    pages, empty = parse_document(path)
    assert empty == [1]  # M2 行为不变：无文本层页登记
    failed = await ocr_fill_pages(path, pages, empty)

    assert failed == []
    assert "扫描页识别出的正文" in pages[0].text
    # 确实走了渲染取图（generate 收到图像）
    assert any(c[0] == "generate" for c in fake.calls)


@pytest.mark.asyncio
async def test_pptx_picture_only_slide_backfilled(tmp_path, monkeypatch):
    """PPT 纯图页：取幻灯片内嵌图片字节喂 OCR，回填后该页有文本。"""
    from pptx import Presentation
    from pptx.util import Inches

    from app.ingest.ocr import ocr_fill_pages
    from app.ingest.parsers import parse_document

    img_path = tmp_path / "chart.png"
    img_path.write_bytes(_png_bytes())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), Inches(2), Inches(2))
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    fake = _FakeOcrGateway(active=[], reply="图表标题：实验数据曲线")
    monkeypatch.setattr("app.ingest.ocr.gateway", fake)

    pages, empty = parse_document(path)
    assert empty == [1]  # 无文字只有图 → 登记
    failed = await ocr_fill_pages(path, pages, empty)
    assert failed == []
    assert "实验数据曲线" in pages[0].text


# ===== 流水线端到端（图片入库 → 可切块可向量化） =====


@pytest.mark.asyncio
async def test_pipeline_end_to_end_image_ingest(tmp_path, db_session, monkeypatch, fake_embed_fn):
    """图片文件走完整 ingest_file：OCR 回填 → 切块 → 假向量入库 → empty_pages 清空。"""
    from app.db import crud
    from app.ingest.pipeline import ingest_file
    from app.rag import embeddings
    from app.rag.vector_store import ChromaStore
    import chromadb

    # 与 test_kb_chat 同款补丁：store_for 换内存客户端、向量化换假向量
    eph = chromadb.EphemeralClient()

    def fake_store_for(user_id, group_id, client=None):
        return ChromaStore(collection=f"u{user_id}g{group_id}", client=eph)

    monkeypatch.setattr("app.rag.vector_store.store_for", fake_store_for)
    monkeypatch.setattr("app.ingest.pipeline.store_for", fake_store_for)
    monkeypatch.setattr(embeddings, "embed_texts", fake_embed_fn)
    monkeypatch.setattr(
        "app.ingest.ocr.gateway",
        _FakeOcrGateway(active=[], reply="质能方程 E 等于 m c 平方"),
    )

    user = crud.create_user(db_session, username="ocr_user", password_hash="x")
    group = crud.create_kb_group(db_session, user_id=user.id, name="物理")
    path = tmp_path / "formula.png"
    path.write_bytes(_png_bytes())

    result = await ingest_file(
        db_session,
        user_id=user.id,
        group_id=group.id,
        file_name="formula.png",
        file_path=path,
    )
    assert result.skipped_identical is False
    assert result.page_count == 1
    assert result.empty_pages == []  # OCR 成功回填 → 清单清空
    assert result.chunk_count >= 1  # 识别出的文本成功切块入库
    assert result.added == result.chunk_count  # 首轮全量 added

    # 指纹与文档状态一致
    doc = crud.get_document(db_session, user_id=user.id, doc_id=result.doc_id)
    assert doc.status == "ready"
    fps = crud.list_chunk_fingerprints(db_session, document_id=doc.id)
    assert len(fps) == result.chunk_count
