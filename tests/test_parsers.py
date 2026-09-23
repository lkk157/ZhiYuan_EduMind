# -*- coding: utf-8 -*-
"""
解析器单测（M2）：PDF/Word/PPT 统一成「逐页文本 + 无文本层页码」的契约行为。

为什么要手工拼最小 PDF：pypdf 只能读不能写带文字的 PDF，
而「有文本层 / 无文本层」两条路径恰恰是溯源与 M4 OCR 交接的核心语义，
必须用真实 PDF 字节验证 extract_text，mock 掉就等于没测。
"""
import io

import pytest
from pypdf import PdfWriter

from app.core.exceptions import AppError
from app.ingest.parsers import parse_document


def _build_pdf(page_texts: list[str]) -> bytes:
    """手工拼一个极简合法 PDF：每页一个 Tj 文本流（ASCII 文本即可验证抽取）。

    为什么不引入 reportlab 之类的新依赖造测试 PDF：技术栈锁定（CLAUDE.md §4），
    为测试引入生产用不到的渲染库不值得；PDF 字节拼接 30 行搞定且完全确定。
    """
    body = b"%PDF-1.4\n"
    offsets: list[int] = []

    def add_obj(payload: bytes) -> None:
        nonlocal body
        offsets.append(len(body))
        body += f"{len(offsets)} 0 obj\n".encode() + payload + b"\nendobj\n"

    n = len(page_texts)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    add_obj(b"<< /Type /Catalog /Pages 2 0 R >>")
    add_obj(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    add_obj(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(page_texts):
        # PDF 字符串里的反斜杠/圆括号必须转义，否则 xref 后内容流解析错位
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        add_obj(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
            ).encode()
        )
        add_obj(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    xref_pos = len(body)
    # xref 声明的条目数 = 对象数 + 1（0 号空闲条目也算一条），写少了 pypdf 会读出 Null 对象
    count = len(offsets) + 1
    body += f"xref\n0 {count}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        body += f"{off:010d} 00000 n \n".encode()
    body += (
        f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF"
    ).encode()
    return body


def test_pdf_text_pages_and_page_numbers(tmp_path):
    """有文本层 PDF：逐页抽出文本，页码 = 真实页码（1 起），empty_pages 为空。"""
    path = tmp_path / "a.pdf"
    path.write_bytes(_build_pdf(["Hello ZhiYuan", "Page two text"]))
    pages, empty = parse_document(path)
    assert empty == []
    assert [p.page_no for p in pages] == [1, 2]
    assert "Hello ZhiYuan" in pages[0].text
    assert "Page two text" in pages[1].text


def test_pdf_blank_pages_go_to_empty_pages(tmp_path):
    """无文本层页（扫描页）：text 置空串、页码进 empty_pages——不产占位文本（防污染检索）。"""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "scan.pdf"
    with path.open("wb") as f:
        writer.write(f)
    pages, empty = parse_document(path)
    assert empty == [1, 2]  # M4 OCR 按这份清单逐页补识别
    assert all(p.text == "" for p in pages)
    assert [p.page_no for p in pages] == [1, 2]  # 空页保留在列表里，page_count 才是真实页数


def test_docx_logical_pages_keep_paragraphs_intact(tmp_path):
    """DOCX 逻辑页：段落累积约 1000 字符切页、段落绝不拆开、页码 1 起连续。"""
    from docx import Document

    doc = Document()
    para = "知" * 400  # 5 段 × 400 字符：按 1000 字符软目标应切出 3 个逻辑页
    for _ in range(5):
        doc.add_paragraph(para)
    path = tmp_path / "b.docx"
    doc.save(str(path))

    pages, empty = parse_document(path)
    assert empty == []  # 逻辑页天然非空
    assert [p.page_no for p in pages] == [1, 2, 3]
    # 段落完整性：每个 400 字符段落必须完整出现在某一页里（绝不撕开）
    for p in pages:
        for part in p.text.split("\n"):
            assert len(part) in (400, 800)  # 整段原样，或两段拼页
    assert pages[0].text.count("知" * 400) == 2  # 第一页装下两段


def test_docx_single_long_paragraph_not_split(tmp_path):
    """单段超长时整段独占一页（段落不拆是硬约束，「约 1000」只是软目标）。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("梯" * 1500)
    path = tmp_path / "long.docx"
    doc.save(str(path))
    pages, empty = parse_document(path)
    assert len(pages) == 1
    assert pages[0].text == "梯" * 1500


def test_pptx_slide_numbers_and_image_only_slide(tmp_path):
    """PPTX：幻灯片序号=页码（1 起）；纯图/空白页进 empty_pages（留给 M4 OCR）。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank_layout)
    tb = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "gradient descent"
    prs.slides.add_slide(blank_layout)  # 第 2 页：空白（模拟纯图页）
    path = tmp_path / "c.pptx"
    prs.save(str(path))

    pages, empty = parse_document(path)
    assert [p.page_no for p in pages] == [1, 2]
    assert "gradient descent" in pages[0].text
    assert empty == [2]
    assert pages[1].text == ""


def test_unsupported_suffix_rejected(tmp_path):
    """非三种后缀在入口就用人话拒绝（AppError 400），不到解析器深处报天书。"""
    path = tmp_path / "x.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(AppError) as e:
        parse_document(path)
    assert e.value.code == 400
    assert "PDF" in e.value.message
