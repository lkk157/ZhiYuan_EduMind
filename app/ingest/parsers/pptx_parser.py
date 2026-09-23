# -*- coding: utf-8 -*-
"""
PPT 解析：python-pptx 每张幻灯片 = 1 页，抽取文本框 / 占位符与表格单元格文字。

幻灯片序号天然稳定（1 起），溯源锚点直接用序号，不像 Word 需要造逻辑页。
"""
from pathlib import Path

from pptx import Presentation

from app.ingest.parsers import ParsedPage


def parse_pptx(path: Path) -> tuple[list[ParsedPage], list[int]]:
    """每张幻灯片 = 1 页（幻灯片序号 1 起），返回 (全部页, 纯图/图表页码列表)。

    抽取范围按契约只收两类：has_text_frame 的文本框与标题/正文占位符、
    has_table 的表格单元格文字；形状按幻灯片上的原始顺序拼接，输出确定。

    为什么纯图片 / 图表页进 empty_pages 而不塞占位文本：
    与 PDF 空页同一取舍（防污染检索）——图示里的知识只能靠视觉理解，
    占位文本进向量库只会挤掉有效结果；页码记入 empty_pages，
    M4 OCR 接管时按清单逐页识别回填（图表理解再往后放）。
    """
    prs = Presentation(str(path))
    pages: list[ParsedPage] = []
    empty_pages: list[int] = []
    for page_no, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            # 文本框与标题/正文占位符都带 text_frame，原文收集（含框内换行）
            if shape.has_text_frame:
                text = shape.text_frame.text
                if text.strip():
                    parts.append(text)
            # 表格：逐行拼单元格文字，行内用制表符分隔保留「列」的语义
            if shape.has_table:
                for row in shape.table.rows:
                    line = "\t".join(cell.text for cell in row.cells)
                    if line.strip():
                        parts.append(line)
        text = "\n".join(parts)
        if not text.strip():
            empty_pages.append(page_no)
            text = ""  # 空白归一成空串，与 PDF 空页约定一致
        pages.append(ParsedPage(page_no=page_no, text=text))
    return pages, empty_pages
